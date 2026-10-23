"""Trusted bounded document parser. Receives data through stdin, never a host path."""
import base64
import io
import json
import resource
import sys
import warnings
import zipfile

IMAGE_SUFFIXES = {'.png', '.jpg', '.jpeg', '.webp'}
IMAGE_FORMATS = {'PNG', 'JPEG', 'WEBP'}
MAX_SIDE = 1568          # Long side sent to the model; larger images cost tokens without adding much.
MAX_PDF_IMAGE_PAGES = 10  # Scanned PDFs: pages sent as images.


def open_image(data):
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = 40_000_000
    warnings.simplefilter('error', Image.DecompressionBombWarning)
    image = Image.open(io.BytesIO(data))
    image.load()
    return image


def encode(image):
    # convert() returns a plain in-memory image: saving one opened from data would ask for the working
    # directory, which the sandbox denies.
    image = image.convert('L' if image.mode == 'L' else 'RGB')
    image.thumbnail((MAX_SIDE, MAX_SIDE))
    output = io.BytesIO()
    image.save(output, 'JPEG', quality=85)
    return {'mime': 'image/jpeg', 'data': base64.b64encode(output.getvalue()).decode()}


def page_images(reader):
    """Largest embedded image of each of the first pages: what a scanned page is made of."""
    images = []
    for page in reader.pages[:MAX_PDF_IMAGE_PAGES]:
        found = sorted(page.images, key=lambda item: len(item.data), reverse=True)
        if found:
            images.append(found[0].data)
    return images


def model_images(suffix, data):
    if suffix in IMAGE_SUFFIXES:
        return [encode(open_image(data))]
    from pypdf import PdfReader
    return [encode(open_image(image)) for image in page_images(PdfReader(io.BytesIO(data), strict=True))]


def main():
    resource.setrlimit(resource.RLIMIT_CPU, (5, 6))
    if sys.platform != 'darwin':
        resource.setrlimit(resource.RLIMIT_AS, (384 * 1024 * 1024,) * 2)
    payload = json.loads(sys.stdin.buffer.read(3_000_000))
    data = base64.b64decode(payload['data'], validate=True)
    suffix = payload['suffix']
    if payload.get('mode') == 'images':
        try:
            images = model_images(suffix, data)
            if not images:
                raise ValueError('no images')
            print(json.dumps({'images': images}))
        except Exception:
            print(json.dumps({'error': '无法把这个附件转成图片：请检查文件是否损坏或过大。'}))
        return
    if suffix in IMAGE_SUFFIXES:
        try:
            image = open_image(data)
            if image.format not in IMAGE_FORMATS:
                raise ValueError('unsupported format')
            print(json.dumps({'text': f'[图片附件 {image.width}×{image.height}：内容附在用户发送它的那条消息里，read_file 只能看到这句说明]'},
                             ensure_ascii=False))
        except Exception:
            print(json.dumps({'error': '图片无法识别：请上传未损坏的 PNG、JPG 或 WEBP，单张像素不超过 4000 万。'}))
        return
    try:
        if suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ValueError('请上传未加密的 PDF。')
            if len(reader.pages) > 40:
                raise ValueError('当前一次最多解析 40 页，请拆分文件。')
            text = '\n\n'.join((page.extract_text() or '')[:30000] for page in reader.pages)[:120000]
            if not text.strip() and page_images(reader):
                pages = min(len(reader.pages), MAX_PDF_IMAGE_PAGES)
                text = (f'[扫描版 PDF，共 {len(reader.pages)} 页，未提取到文字：前 {pages} 页以图片附在用户发送它的那条消息里，'
                        'read_file 只能看到这句说明]')
        elif suffix == '.docx':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20_000_000 or len(archive.infolist()) > 1500:
                    raise ValueError('DOCX 解压内容过大。')
            from docx import Document
            doc = Document(io.BytesIO(data))
            text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])[:120000]
        else:
            text = data.decode('utf-8-sig')[:120000]
        if not text.strip():
            raise ValueError('未提取到文字。扫描 PDF 或图片需先进行 OCR。')
        print(json.dumps({'text': text}, ensure_ascii=False))
    except Exception:
        print(json.dumps({'error': '文件解析失败：请检查格式、加密状态及页数；扫描件需先进行 OCR。'}))


if __name__ == '__main__':
    main()
